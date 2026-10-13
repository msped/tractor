"""
Training and inference must read a DOCX identically.

Two separate implementations of "turn a DOCX into NER text" used to exist —
one in the training task, one in the extraction module — and they disagreed on
every real document: over blank paragraphs, and over hyperlinked text that the
training side dropped entirely. A model trained on a string the pipeline never
produces is being asked, at inference, to label something it has not seen.
"""

import os
import tempfile

from django.test import TestCase
from docx import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement

from training.extraction import (
    extract_annotated_text,
    extract_document_structure,
)
from training.tasks import HIGHLIGHT_COLOR_TO_LABEL, _trim_spans

from .base import NetworkBlockerMixin


def _highlighted(paragraph, text, colour):
    run = paragraph.add_run(text)
    run.font.highlight_color = colour
    return run


class ExtractionParityTests(NetworkBlockerMixin, TestCase):
    def _document(self):
        """A DOCX exercising every construct the two paths used to differ on."""
        doc = DocxDocument()

        doc.add_paragraph("Occurrence details")
        doc.add_paragraph("")  # blank: kept by one path, dropped by the other

        para = doc.add_paragraph("Reported by ")
        _highlighted(para, "BLOGGS, Joe", WD_COLOR_INDEX.BRIGHT_GREEN)
        para.add_run(" on duty")

        doc.add_paragraph("   ")  # whitespace-only

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Report no.:"
        table.cell(0, 1).text = "23123456789"
        cell = table.cell(1, 0)
        cell.text = ""
        _highlighted(
            cell.paragraphs[0], "Involved officers:", WD_COLOR_INDEX.TURQUOISE
        )
        table.cell(1, 1).text = "#0712345"

        return doc

    def _write(self, doc):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            path = f.name
        doc.save(path)
        return path

    def test_training_and_inference_produce_identical_text(self):
        path = self._write(self._document())
        try:
            _, _, inference_text = extract_document_structure(path)
            training_text, _ = extract_annotated_text(
                path, HIGHLIGHT_COLOR_TO_LABEL
            )

            self.assertEqual(training_text, inference_text)
        finally:
            os.remove(path)

    def test_highlight_offsets_point_at_the_highlighted_text(self):
        """
        The offsets are only meaningful if they land on the right characters —
        an off-by-one here teaches the model the wrong span silently.
        """
        path = self._write(self._document())
        try:
            text, spans = extract_annotated_text(
                path, HIGHLIGHT_COLOR_TO_LABEL
            )
            found = {(text[s:e], label) for s, e, label in spans}

            self.assertIn(("BLOGGS, Joe", "THIRD_PARTY"), found)
            self.assertIn(("Involved officers:", "OPERATIONAL"), found)
        finally:
            os.remove(path)

    def test_hyperlinked_text_is_not_dropped(self):
        """
        ``paragraph.runs`` omits runs nested in a hyperlink, so walking it
        loses that text — and any highlight on it — without any error.
        """
        doc = DocxDocument()
        para = doc.add_paragraph("See ")
        para.add_run("plain text ")
        linked = _highlighted(
            para, "LION HOTEL, WINSFORD", WD_COLOR_INDEX.BRIGHT_GREEN
        )
        # python-docx cannot author a hyperlink, so nest the run by hand —
        # this is the shape that hides text from ``paragraph.runs``.
        hyperlink = OxmlElement("w:hyperlink")
        linked._element.addprevious(hyperlink)
        hyperlink.append(linked._element)

        self.assertEqual(
            [r.text for r in para.runs], ["See ", "plain text "]
        )  # the hyperlinked run is invisible to paragraph.runs

        path = self._write(doc)
        try:
            text, spans = extract_annotated_text(
                path, HIGHLIGHT_COLOR_TO_LABEL
            )

            self.assertIn("LION HOTEL, WINSFORD", text)
            self.assertEqual(
                [(text[s:e], label) for s, e, label in spans],
                [("LION HOTEL, WINSFORD", "THIRD_PARTY")],
            )
        finally:
            os.remove(path)

    def test_adjacent_runs_of_one_label_merge(self):
        """Formatting splits a phrase into runs; the annotation is one span."""
        doc = DocxDocument()
        para = doc.add_paragraph()
        _highlighted(para, "Involved ", WD_COLOR_INDEX.TURQUOISE)
        bold = _highlighted(para, "officers", WD_COLOR_INDEX.TURQUOISE)
        bold.bold = True
        _highlighted(para, ":", WD_COLOR_INDEX.TURQUOISE)

        path = self._write(doc)
        try:
            text, spans = extract_annotated_text(
                path, HIGHLIGHT_COLOR_TO_LABEL
            )
            self.assertEqual(len(spans), 1)
            start, end, label = spans[0]
            self.assertEqual(text[start:end], "Involved officers:")
            self.assertEqual(label, "OPERATIONAL")
        finally:
            os.remove(path)

    def test_unmapped_highlight_colour_is_ignored(self):
        doc = DocxDocument()
        para = doc.add_paragraph()
        _highlighted(para, "Not an annotation", WD_COLOR_INDEX.YELLOW)

        path = self._write(doc)
        try:
            _, spans = extract_annotated_text(path, HIGHLIGHT_COLOR_TO_LABEL)
            self.assertEqual(spans, [])
        finally:
            os.remove(path)

    def test_non_docx_returns_nothing(self):
        text, spans = extract_annotated_text(
            "/some/file.pdf", HIGHLIGHT_COLOR_TO_LABEL
        )
        self.assertIsNone(text)
        self.assertIsNone(spans)


class TrimSpansTests(TestCase):
    """Highlights routinely overrun the word by a space."""

    def test_edges_are_trimmed(self):
        text = "Reported by BLOGGS, Joe on duty"
        self.assertEqual(
            _trim_spans([(11, 24, "THIRD_PARTY")], text),
            [(12, 23, "THIRD_PARTY")],
        )

    def test_whitespace_only_span_is_discarded(self):
        self.assertEqual(_trim_spans([(0, 3, "OPERATIONAL")], "   text"), [])

    def test_already_tight_span_is_unchanged(self):
        text = "BLOGGS, Joe"
        self.assertEqual(
            _trim_spans([(0, 11, "THIRD_PARTY")], text),
            [(0, 11, "THIRD_PARTY")],
        )
