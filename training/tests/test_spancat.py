"""
Tests for the training core that does not need Django.

Everything here runs without the database, which is the property that lets
the sweep script train the same way the application does. Nothing in this
file may reach for a model, a setting or a connection.
"""

import tempfile
from collections import Counter
from pathlib import Path
from unittest.mock import MagicMock, patch

from django.test import SimpleTestCase
from docx import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX

from training.spancat import (
    BASE_MODEL,
    CorpusStats,
    EpochRecord,
    TrainingStopped,
    base_model_present,
    ensure_base_model,
    load_annotated_docs,
    prepare_examples,
    run_training_loop,
    split_long_span,
)
from training.tests.base import NetworkBlockerMixin


def write_docx(path, runs):
    """Write a DOCX whose runs are (text, colour-or-None) pairs."""
    doc = DocxDocument()
    para = doc.add_paragraph()
    for text, colour in runs:
        run = para.add_run(text)
        if colour:
            run.font.highlight_color = colour
    doc.save(path)


class CorpusStatsTests(SimpleTestCase):
    def test_kept_is_what_survived_alignment(self):
        stats = CorpusStats(total=10, dropped=3)

        self.assertEqual(stats.kept, 7)

    def test_over_counts_annotations_longer_than_the_ceiling(self):
        stats = CorpusStats(
            total=6, token_lengths=Counter({5: 2, 20: 3, 50: 1})
        )

        self.assertEqual(stats.over(20), 1)
        self.assertEqual(stats.over(4), 6)
        self.assertEqual(stats.over(50), 0)

    def test_unreachable_share_is_the_fraction_a_ceiling_cannot_propose(self):
        """The number that decides one ceiling over another."""
        stats = CorpusStats(total=4, token_lengths=Counter({10: 3, 90: 1}))

        self.assertEqual(stats.unreachable_share(40), 0.25)
        self.assertEqual(stats.unreachable_share(90), 0.0)

    def test_unreachable_share_of_an_empty_corpus_is_zero(self):
        self.assertEqual(CorpusStats().unreachable_share(40), 0.0)

    def test_token_lengths_do_not_depend_on_the_ceiling(self):
        """
        One pass answers the coverage question for every ceiling, which is
        what makes the corpus survey cheap enough to run before training.
        """
        import spacy

        nlp = spacy.blank("en")
        text = " ".join(f"word{i}" for i in range(30))
        data = [(text, {"entities": [(0, len(text), "OPERATIONAL")]})]

        _, loose = prepare_examples(nlp, data, max_tokens=None)
        _, tight = prepare_examples(nlp, data, max_tokens=5)

        self.assertEqual(loose.token_lengths, tight.token_lengths)
        self.assertEqual(loose.unreachable_share(5), 1.0)
        self.assertEqual(loose.unreachable_share(30), 0.0)


class EpochRecordTests(SimpleTestCase):
    def test_reads_as_the_progress_line_a_run_prints(self):
        record = EpochRecord(
            epoch=2,
            loss=12.5,
            precision=0.5,
            recall=0.25,
            f1=0.333,
            seconds=9.0,
            improved=True,
        )

        self.assertEqual(
            str(record),
            "Epoch 2/15 - Loss: 12.5000 - Dev F1: 0.3330 P: 0.5000 R: 0.2500",
        )


class BaseModelTests(NetworkBlockerMixin, SimpleTestCase):
    """The base model has to be there before a run can start."""

    def test_present_when_the_package_is_installed(self):
        with patch("spacy.util.is_package", return_value=True):
            self.assertTrue(base_model_present())

    def test_an_installed_model_is_not_downloaded_again(self):
        with patch("spacy.util.is_package", return_value=True):
            with patch("spacy.cli.download") as download:
                downloaded = ensure_base_model()

        self.assertFalse(downloaded)
        download.assert_not_called()

    def test_a_missing_model_is_fetched(self):
        with patch("spacy.util.is_package", side_effect=[False, True]):
            with patch("spacy.cli.download") as download:
                downloaded = ensure_base_model(log=lambda _msg: None)

        self.assertTrue(downloaded)
        download.assert_called_once_with(BASE_MODEL)

    def test_refuses_to_download_when_told_not_to(self):
        with patch("spacy.util.is_package", return_value=False):
            with patch("spacy.cli.download") as download:
                with self.assertRaises(RuntimeError) as caught:
                    ensure_base_model(download=False)

        download.assert_not_called()
        self.assertIn(BASE_MODEL, str(caught.exception))

    def test_raises_when_a_download_leaves_it_still_missing(self):
        """
        Installing into a different environment than the one running is the
        usual cause, and silently training nothing would be worse.
        """
        with patch("spacy.util.is_package", return_value=False):
            with patch("spacy.cli.download"):
                with self.assertRaises(RuntimeError):
                    ensure_base_model(log=lambda _msg: None)


class LoadAnnotatedDocsTests(SimpleTestCase):
    """Reading a corpus from files alone, with no TrainingDocument rows."""

    def test_highlights_become_labelled_entities(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.docx"
            write_docx(
                path,
                [
                    ("Reported by ", None),
                    ("SMITH, J.", WD_COLOR_INDEX.BRIGHT_GREEN),
                    (" and logged as ", None),
                    ("crime ref 12/345", WD_COLOR_INDEX.TURQUOISE),
                ],
            )

            train_data, skipped = load_annotated_docs([path])

        self.assertEqual(skipped, [])
        text, annotations = train_data[0]
        labelled = {
            text[start:end]: label
            for start, end, label in annotations["entities"]
        }
        self.assertEqual(
            labelled,
            {"SMITH, J.": "THIRD_PARTY", "crime ref 12/345": "OPERATIONAL"},
        )

    def test_a_document_with_no_highlights_is_skipped_with_a_reason(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "plain.docx"
            write_docx(path, [("Nothing is marked up in here.", None)])

            train_data, skipped = load_annotated_docs([path])

        self.assertEqual(train_data, [])
        self.assertEqual(skipped, [(str(path), "no highlighted spans")])

    def test_a_file_that_is_not_a_docx_is_skipped(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "notes.docx"
            path.write_text("this is not a word document")

            train_data, skipped = load_annotated_docs([path])

        self.assertEqual(train_data, [])
        self.assertEqual(skipped, [(str(path), "no text extracted")])

    def test_an_unreadable_document_does_not_abandon_the_rest(self):
        with patch(
            "training.spancat.extract_annotated_text",
            side_effect=OSError("disk went away"),
        ):
            train_data, skipped = load_annotated_docs(["/some/report.docx"])

        self.assertEqual(train_data, [])
        self.assertEqual(skipped, [("/some/report.docx", "disk went away")])


class PrepareExamplesCoreTests(SimpleTestCase):
    def test_counts_what_it_kept_dropped_and_divided(self):
        import spacy

        nlp = spacy.blank("en")
        long_text = " ".join(f"word{i}" for i in range(12))
        data = [
            (long_text, {"entities": [(0, len(long_text), "OPERATIONAL")]}),
            ("Jane Doe called", {"entities": [(0, 8, "THIRD_PARTY")]}),
        ]

        examples, stats = prepare_examples(nlp, data, max_tokens=5)

        self.assertEqual(len(examples), 2)
        self.assertEqual(stats.total, 2)
        self.assertEqual(stats.dropped, 0)
        self.assertEqual(stats.split, 1)
        self.assertGreater(stats.pieces, 1)

    def test_nothing_interrupts_a_run_by_default(self):
        """
        The application passes a stop check; a standalone run has none, and
        must not need one.
        """
        import spacy

        nlp = spacy.blank("en")
        data = [(f"text {i}", {"entities": []}) for i in range(40)]

        examples, _stats = prepare_examples(nlp, data)

        self.assertEqual(len(examples), 40)

    def test_the_stop_check_is_offered_the_stage_it_reached(self):
        import spacy

        nlp = spacy.blank("en")
        data = [(f"text {i}", {"entities": []}) for i in range(25)]
        stages = []

        prepare_examples(nlp, data, should_stop=stages.append)

        self.assertEqual(
            stages,
            ["while preparing example 10", "while preparing example 20"],
        )

    def test_a_raising_stop_check_abandons_the_phase(self):
        import spacy

        nlp = spacy.blank("en")
        data = [(f"text {i}", {"entities": []}) for i in range(40)]

        def stop(stage):
            raise TrainingStopped(stage)

        with self.assertRaises(TrainingStopped):
            prepare_examples(nlp, data, should_stop=stop)


class SplitLongSpanCoreTests(SimpleTestCase):
    def test_pieces_cover_the_same_text_as_the_whole(self):
        import spacy

        nlp = spacy.blank("en")
        doc = nlp("alpha bravo charlie; delta echo foxtrot golf hotel india")
        span = doc.char_span(0, len(doc.text), label="OPERATIONAL")

        pieces = split_long_span(span, "OPERATIONAL", 4)

        self.assertGreater(len(pieces), 1)
        self.assertEqual(pieces[0].start_char, span.start_char)
        self.assertEqual(pieces[-1].end_char, span.end_char)
        self.assertTrue(all(len(p) <= 4 for p in pieces))


class RunTrainingLoopCoreTests(SimpleTestCase):
    def _nlp(self, f1=0.5):
        nlp = MagicMock()
        nlp.resume_training.return_value = MagicMock()
        nlp.evaluate.return_value = {
            "spans_sc_f": f1,
            "spans_sc_p": f1,
            "spans_sc_r": f1,
        }
        return nlp

    def test_history_records_every_epoch_that_ran(self):
        nlp = self._nlp()

        with tempfile.TemporaryDirectory() as out:
            best, history = run_training_loop(nlp, [], [], out)

        # One improvement, then patience exhausted on identical scores.
        self.assertEqual([r.epoch for r in history], [1, 2, 3, 4])
        self.assertEqual(
            [r.improved for r in history], [True, False, False, False]
        )
        self.assertEqual(best["spans_sc_f"], 0.5)
        self.assertTrue(all(r.seconds >= 0 for r in history))

    def test_progress_is_reported_as_it_goes(self):
        nlp = self._nlp()
        seen = []

        with tempfile.TemporaryDirectory() as out:
            run_training_loop(nlp, [], [], out, on_epoch=seen.append)

        self.assertEqual([type(item) for item in seen[:4]], [EpochRecord] * 4)
        self.assertIn("Early stopping", seen[-1])

    def test_a_run_nobody_can_stop_trains_to_the_end(self):
        nlp = self._nlp()

        with tempfile.TemporaryDirectory() as out:
            _best, history = run_training_loop(nlp, [], [], out)

        self.assertTrue(nlp.evaluate.called)
        self.assertTrue(history)

    def test_every_batch_is_offered_the_stop_check(self):
        """
        Checked per batch rather than per epoch: at a high ceiling a single
        batch can run for a long time, so a stop honoured only between epochs
        would appear to be ignored.
        """
        nlp = self._nlp()
        stages = []

        with tempfile.TemporaryDirectory() as out:
            run_training_loop(
                nlp,
                [MagicMock(), MagicMock()],
                [],
                out,
                should_stop=stages.append,
            )

        self.assertTrue(nlp.update.called)
        self.assertIn("during epoch 1", stages)
        self.assertIn("at the start of epoch 1", stages)

    def test_a_raising_stop_check_abandons_the_loop(self):
        nlp = self._nlp()

        def stop(stage):
            raise TrainingStopped(stage)

        with tempfile.TemporaryDirectory() as out:
            with self.assertRaises(TrainingStopped):
                run_training_loop(nlp, [], [], out, should_stop=stop)

        nlp.update.assert_not_called()
        nlp.evaluate.assert_not_called()

    def test_the_final_state_is_saved_when_nothing_ever_improved(self):
        """A model scoring zero is still the model this corpus produced."""
        nlp = self._nlp(f1=0.0)

        with tempfile.TemporaryDirectory() as out:
            best, _history = run_training_loop(nlp, [], [], out)

        nlp.to_disk.assert_called_with(out)
        self.assertEqual(best["spans_sc_f"], 0.0)
